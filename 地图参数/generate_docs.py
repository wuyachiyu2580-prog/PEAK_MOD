#!/usr/bin/env python3
"""从聚合中间数据生成 HTML 网页 + Excel 电子表格 (中英双语版)"""
"""v3: +物体对比 +自定义精简参数"""

import json, os, sys
from datetime import datetime
from pathlib import Path

BASE = Path(r"c:\Users\Administrator\Desktop\MOD\PEAK")
OUT_DIR = BASE / "地图参数"
INTERMEDIATE = OUT_DIR / "_aggregated_data.json"

sys.path.insert(0, str(OUT_DIR))
from build_docs import TRANSLATIONS

SEGMENT_INFO = {
    "Beach_Segment": {"order":1,"enCN":"海滩 (Beach)","desc":"第一关·热带海滩。棕榈树、椰子和水母是主要特色，部分变体有蛇、黑沙或大量水母。"},
    "Jungle_Segment": {"order":2,"enCN":"雨林 (Jungle)","desc":"第二关·热带雨林。茂密植被、藤蔓和炸弹是主要特色，变体差异极大——从岩浆到空中平台都有。"},
    "Roots Segment": {"order":2,"enCN":"森蕈 (Roots)","desc":"第二关·森蕈森林。巨型蘑菇、红杉和深水是主要特色，洞穴、甲虫和秃林等变体变化丰富。"},
    "Snow_Segment": {"order":3,"enCN":"雪山 (Snow)","desc":"第三关·雪山。冰雪覆盖的峭壁，有岩浆、尖刺和喷泉等危险变体。"},
    "Desert_Segment": {"order":3,"enCN":"沙漠 (Desert)","desc":"第三关·沙漠。仙人掌、蝎子、炸药和龙卷风是主要特色，变体差异明显。"},
    "Caldera_Segment": {"order":4,"enCN":"破火山口 (Caldera)","desc":"第四关·破火山口。环形火山地形，有熔岩河和瀑布，无变体系统。"},
    "Volcano_Segment": {"order":5,"enCN":"火山口 (Volcano)","desc":"第五关·火山口。最终关卡，火山内部，有岩浆和上升熔岩，无变体系统。"},
}
SEGMENT_ORDER=["Beach_Segment","Jungle_Segment","Roots Segment","Snow_Segment","Desert_Segment","Caldera_Segment","Volcano_Segment"]

def fix_variant_name(vn):
    if vn is None or vn=="null" or vn=="": return "默认"
    return vn

def load_data():
    with open(str(INTERMEDIATE),"r",encoding="utf-8") as f: return json.load(f)

def escape_html(s):
    return str(s).replace("&","&amp;").replace("<","&lt;").replace(">","&gt;").replace('"',"&quot;")

def format_val_short(v):
    if v is None: return "—"
    if isinstance(v,bool): return "✓" if v else "✗"
    if isinstance(v,(int,float)):
        if isinstance(v,float):
            if abs(v)<0.0001: return "0"
            if abs(v-round(v))<0.001: return str(int(v))
            return f"{v:.1f}"
        return str(v)
    if isinstance(v,dict):
        if "x" in v:
            z=v.get("z",None)
            return f"({v['x']:.0f},{v['y']:.0f},{z:.0f})" if z is not None else f"({v['x']:.0f},{v['y']:.0f})"
        return "对象"
    return str(v)[:40]

def _bilingual(cn,en):
    if not cn or cn==en: return escape_html(en)
    if not en: return escape_html(cn)
    return f'{escape_html(cn)} <span class="en-name">({escape_html(en)})</span>'

def _tr(key):
    return TRANSLATIONS.get(key,key)

def _is_mute_row(pk,pv):
    return pk=="mute" and pv==False

def _build_step_html(st):
    rows=[]
    for pk,pv in st.get("properties",{}).items():
        cn=_tr(pk)
        cls=' class="mute-row"' if _is_mute_row(pk,pv) else ""
        rows.append(f'<tr{cls} data-param="{escape_html(pk)}"><td class="prop-name">{_bilingual(cn,pk)}</td><td class="prop-val">{escape_html(format_val_short(pv))}</td><td class="prop-cat">属性</td></tr>')
    for mname,mprops in st.get("modifiers",{}).items():
        mcn=_tr(mname)
        for pk,pv in mprops.items():
            cn=_tr(pk)
            cls=' class="mute-row"' if _is_mute_row(pk,pv) else ""
            rows.append(f'<tr{cls} data-param="{escape_html(pk)}"><td class="prop-name">{_bilingual(cn,pk)}</td><td class="prop-val">{escape_html(format_val_short(pv))}</td><td class="prop-cat mod">{_bilingual(mcn,mname)}</td></tr>')
    for cname,cprops in st.get("constraints",{}).items():
        ccn=_tr(cname)
        for pk,pv in cprops.items():
            cn=_tr(pk)
            cls=' class="mute-row"' if _is_mute_row(pk,pv) else ""
            rows.append(f'<tr{cls} data-param="{escape_html(pk)}"><td class="prop-name">{_bilingual(cn,pk)}</td><td class="prop-val">{escape_html(format_val_short(pv))}</td><td class="prop-cat con">{_bilingual(ccn,cname)}</td></tr>')
    return "\n".join(rows)

def gen_html(data):
    segs=data["segments"]
    total_variants=sum(len(s["variants"]) for s in segs.values())
    total_steps=sum(sum(sum(len(g["steps"]) for g in v["groupers"]) for v in s["variants"].values()) for s in segs.values())

    tab_buttons=[]; tab_contents=[]; all_params=set()

    for seg_name in SEGMENT_ORDER:
        if seg_name not in segs: continue
        seg=segs[seg_name]
        info=SEGMENT_INFO.get(seg_name,{"order":0,"enCN":seg_name,"desc":""})
        seg_id=seg_name.replace(" ","_").replace(".","")

        tab_buttons.append(f'<button class="tab-btn" onclick="showTab(\'{seg_id}\')">{info["order"]}. {info["enCN"]}</button>')

        html=f'<div class="tab-content" id="tab-{seg_id}">'
        html+=f'<h2>{info["order"]}. {info["enCN"]} <small>({seg_name})</small></h2>'
        html+=f'<p class="seg-desc">{info["desc"]}</p>'
        v_count=len(seg["variants"])
        html+=f'''<div class="action-bar">
<button onclick="toggleAll('{seg_id}',true)">📂 全部展开</button>
<button onclick="toggleAll('{seg_id}',false)">📁 全部折叠</button>
<span class="sep">|</span>
<button id="btnMute_{seg_id}" onclick="toggleMuteRows('{seg_id}')">🔇 显示静音参数</button>
<span class="sep">|</span>
<button onclick="openSettings('{seg_id}')" title="自定义要显示的参数">⚙ 参数设置</button>
<span class="sep">|</span><span class="stat-inline">{v_count}变体</span></div>'''

        variants=[(vn,v) for vn,v in seg["variants"].items()]
        # 找到默认变体作为基准
        default_v=None
        for vn,v in variants:
            vn_f=fix_variant_name(vn)
            if vn_f in ('默认','Default','default','null',''): default_v=v; break
        if default_v is None and variants: default_v=variants[0][1]
        # 提取基准的步骤集和分组器集
        def _step_set(v):
            s=set()
            for g in v["groupers"]:
                for st in g["steps"]:
                    cn=st["nameCN"];
                    if cn!=st["name"]: s.add(cn)
            return s
        default_steps=_step_set(default_v) if default_v else set()
        default_groupers=set(g["nameCN"] for g in default_v["groupers"]) if default_v else set()
        # 建立每个变体 step_nameCN → (grouper_name, step_name) 映射，供差异链接用
        variant_step_lookup={}
        for vn,v in variants:
            lookup={}
            for g in v["groupers"]:
                for st in g["steps"]:
                    cn=st["nameCN"]
                    if cn!=st["name"]: lookup[cn]=(g["name"], st["name"])
            variant_step_lookup[vn]=lookup

        html+='<div class="quick-table-wrap"><table class="quick-table"><thead><tr>'
        html+='<th>变体</th><th>类型</th><th>样本数</th><th>分组器差异</th><th>步骤</th><th>物件差异 vs 默认</th>'
        html+='</tr></thead><tbody>'
        for vn,v in variants:
            vn_fixed=fix_variant_name(vn)
            v_cn=v.get("nameCN",vn_fixed)
            if v_cn=="None" or v_cn is None: v_cn=vn_fixed
            v_label=_bilingual(v_cn,vn_fixed) if v_cn!=vn_fixed else escape_html(v_cn)
            s_count=sum(len(g["steps"]) for g in v["groupers"])
            is_default=(vn_fixed in ('默认','Default','default','null','') or (default_v is None and vn==variants[0][0]))
            if is_default:
                g_diff='— 基准 —'
                s_diff='— 基准 —'
            else:
                v_groupers=set(g["nameCN"] for g in v["groupers"])
                v_steps=_step_set(v)
                g_parts=[]
                added_g=v_groupers-default_groupers; removed_g=default_groupers-v_groupers
                if added_g: g_parts.append(f'<span class="diff-add">+{",".join(sorted(added_g))}</span>')
                if removed_g: g_parts.append(f'<span class="diff-del">-{",".join(sorted(removed_g))}</span>')
                g_diff=' '.join(g_parts) if g_parts else '同默认'
                s_parts=[]
                added_s=v_steps-default_steps; removed_s=default_steps-v_steps
                lookup=variant_step_lookup.get(vn,{})
                def _anchor(nameCN):
                    if nameCN in lookup:
                        gn,sn=lookup[nameCN]
                        gs=gn.replace(' ','_').replace('(','').replace(')','').replace('.','')
                        ss=sn.replace(' ','_').replace('(','').replace(')','').replace('.','')
                        return f'step-{seg_id}-{gs}-{ss}'
                    return None
                def _dlink(nameCN, cls):
                    a=_anchor(nameCN)
                    if a:
                        return f'<a href="#{a}" class="step-link {cls}" onclick="jumpStep(event,\'{a}\')">{escape_html(nameCN)}</a>'
                    return f'<span class="{cls}">{escape_html(nameCN)}</span>'
                if added_s: s_parts.append('<span class="diff-add">+'+','.join(_dlink(x,'diff-add') for x in sorted(added_s)[:6])+'</span>')
                if removed_s: s_parts.append('<span class="diff-del">-'+','.join(_dlink(x,'diff-del') for x in sorted(removed_s)[:6])+'</span>')
                s_diff=' '.join(s_parts) if s_parts else '同默认'
            html+=f'<tr><td class="vn">{v_label}</td><td>{v["type"]}</td><td>{v["sampleCount"]}</td><td class="g-diff">{g_diff}</td><td>{s_count}</td><td class="special">{s_diff}</td></tr>'
        html+='</tbody></table></div>'

        for vn,v in variants:
            vn_fixed=fix_variant_name(vn)
            v_cn=v.get("nameCN",vn_fixed)
            if v_cn=="None" or v_cn is None: v_cn=vn_fixed
            v_label=f"{v_cn} ({vn_fixed})" if v_cn!=vn_fixed else v_cn
            html+=f'<details class="variant-detail"><summary><strong>{escape_html(v_label)}</strong> — {v["type"]} · {v["sampleCount"]}样本 · {sum(len(g["steps"]) for g in v["groupers"])}步骤</summary>'
            for g in v["groupers"]:
                html+=f'<div class="grouper-section"><h4>{_bilingual(g["nameCN"],g["name"])}</h4>'
                for st in g["steps"]:
                    stype_cn=st.get("typeCN",st["type"])
                    stype_label=f"{stype_cn} ({st['type']})" if stype_cn!=st["type"] else st["type"]
                    st_labels=f'{_bilingual(st["nameCN"],st["name"])} <span class="step-type">{escape_html(stype_label)}</span>'
                    props=st.get("properties",{})
                    nr=props.get("nrOfSpawns",None)
                    chance=props.get("chanceToUseSpawner",None)
                    pp=[]
                    if nr is not None: pp.append(f"生成:{format_val_short(nr)}")
                    if chance is not None:
                        cs=format_val_short(chance)
                        if cs not in ("✓","✗") or cs=="✓": pp.append(f"概率:{cs}")
                    pv_text=" · ".join(pp) if pp else ""
                    preview=f' <span class="preview">{pv_text}</span>' if pv_text else ""
                    vip=f'{escape_html(st["nameCN"])}|{escape_html(v_cn)}|{escape_html(g["nameCN"])}'
                    gs=g["name"].replace(' ','_').replace('(','').replace(')','').replace('.','')
                    ss=st["name"].replace(' ','_').replace('(','').replace(')','').replace('.','')
                    html+=f'<details class="step-detail" id="step-{seg_id}-{gs}-{ss}"><summary><span class="step-label">{st_labels}{preview}</span> <button class="cmp-btn" onclick="addToCompare(this);event.stopPropagation()" title="加入对比" data-stinfo="{vip}">⊕</button></summary>'
                    html+='<table class="prop-table"><thead><tr><th>参数名</th><th>值</th><th>来源</th></tr></thead><tbody>'
                    html+=_build_step_html(st)
                    html+='</tbody></table></details>'
                html+='</div>'
            html+='</details>'
        html+='</div>'
        tab_contents.append(html)

    # 收集所有参数名
    for seg_name in SEGMENT_ORDER:
        if seg_name not in segs: continue
        for v in segs[seg_name]["variants"].values():
            for g in v["groupers"]:
                for st in g["steps"]:
                    for pk in st.get("properties",{}): all_params.add(pk)
                    for m in st.get("modifiers",{}):
                        for pk in st["modifiers"][m]: all_params.add(pk)
                    for c in st.get("constraints",{}):
                        for pk in st["constraints"][c]: all_params.add(pk)

    CORE_PARAMS = ["area","nrOfSpawns","chanceToUseSpawner","layerType"]
    param_checkboxes = ""
    for pk in sorted(all_params):
        cn=_tr(pk); label=f"{cn} ({pk})" if cn!=pk else pk
        checked=' checked' if pk in CORE_PARAMS else ''
        param_checkboxes+=f'<label class="param-cb"><input type="checkbox" value="{escape_html(pk)}"{checked}> {escape_html(label)}</label>\n'

    now=datetime.now().strftime("%Y-%m-%d %H:%M")
    full=f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>PEAK 地图参数手册</title>
<style>
*{{margin:0;padding:0;box-sizing:border-box;}}
body{{font-family:"Microsoft YaHei","PingFang SC","Noto Sans SC",sans-serif;background:#1a1a2e;color:#e0e0e0;line-height:1.6;}}
.header{{background:linear-gradient(135deg,#16213e,#0f3460);padding:30px 20px;text-align:center;border-bottom:3px solid #e94560;}}
.header h1{{font-size:2em;color:#e94560;}}
.header p{{color:#aaa;margin-top:8px;}}
.stats{{display:flex;justify-content:center;gap:30px;margin-top:15px;flex-wrap:wrap;}}
.stat{{background:rgba(255,255,255,0.05);padding:8px 20px;border-radius:20px;font-size:0.9em;}}
.stat strong{{color:#e94560;}}
.stat-inline{{color:#888;font-size:0.85em;}}
.search-bar{{text-align:center;padding:12px 10px;background:#16213e;display:flex;justify-content:center;align-items:center;gap:8px;}}
.search-input-group{{display:flex;align-items:center;}}
.search-bar input{{width:500px;max-width:70%;padding:10px 16px;border:1px solid #444;border-radius:20px 0 0 20px;background:#0f3460;color:#e0e0e0;font-size:1em;outline:none;}}
.search-bar input:focus{{border-color:#e94560;}}
.search-bar .search-clear{{padding:10px 18px;border:1px solid #e94560;border-left:none;background:#e94560;color:#fff;cursor:pointer;font-size:1em;border-radius:0 20px 20px 0;transition:all 0.2s;}}
.search-bar .search-clear:hover{{background:#ff5a7a;}}
.search-bar .search-info{{color:#888;font-size:0.85em;margin-left:8px;white-space:nowrap;}}
.tabs{{display:flex;flex-wrap:wrap;gap:4px;padding:12px 10px;background:#16213e;position:sticky;top:0;z-index:100;justify-content:center;}}
.tab-btn{{padding:10px 18px;border:none;background:#0f3460;color:#ccc;cursor:pointer;border-radius:6px 6px 0 0;font-size:0.95em;transition:all 0.2s;}}
.tab-btn:hover{{background:#e94560;color:#fff;}}
.tab-btn.active{{background:#e94560;color:#fff;font-weight:bold;}}
.content{{max-width:1500px;margin:0 auto;padding:20px;}}
.tab-content{{display:none;}}
.tab-content.active{{display:block;}}
.tab-content h2{{color:#e94560;margin-bottom:5px;}}
.tab-content h2 small{{font-size:0.5em;color:#888;font-weight:normal;}}
.seg-desc{{color:#aaa;margin-bottom:12px;padding:12px 16px;background:rgba(255,255,255,0.03);border-left:3px solid #e94560;border-radius:0 8px 8px 0;}}
.action-bar{{display:flex;gap:8px;margin-bottom:16px;flex-wrap:wrap;align-items:center;}}
.action-bar button{{padding:6px 14px;border:1px solid #555;background:#0f3460;color:#ccc;cursor:pointer;border-radius:4px;font-size:0.85em;transition:all 0.2s;}}
.action-bar button:hover{{border-color:#e94560;color:#fff;}}
.action-bar button.mute-visible{{background:#4ecca3;border-color:#4ecca3;color:#1a1a2e;}}
.action-bar .sep{{color:#555;margin:0 2px;}}
.quick-table-wrap{{overflow-x:auto;margin-bottom:30px;border-radius:8px;}}
.quick-table{{width:100%;border-collapse:collapse;font-size:0.9em;table-layout:fixed;}}
.quick-table th{{background:#0f3460;padding:10px 12px;text-align:left;font-weight:bold;}}
.quick-table td{{padding:8px 12px;border-bottom:1px solid #333;overflow:hidden;text-overflow:ellipsis;white-space:nowrap;}}
.quick-table td:hover{{overflow:visible;white-space:normal;background:#1a1a2e;position:relative;z-index:10;box-shadow:0 2px 8px rgba(0,0,0,0.5);}}
.quick-table tr:hover{{background:rgba(233,69,96,0.08);}}
.quick-table .vn{{color:#e94560;font-weight:bold;}}
.quick-table .special{{color:#aaa;font-size:0.85em;}}
.quick-table .g-diff{{font-size:0.85em;}}
.diff-add{{color:#4ecca3;}}
.diff-del{{color:#e94560;text-decoration:line-through;}}
.step-link{{text-decoration:none;cursor:pointer;}}
.step-link:hover{{text-decoration:underline;}}
.step-link.diff-add{{color:#4ecca3;}}
.step-link.diff-del{{color:#e94560;text-decoration:line-through;}}
.step-link.diff-del:hover{{text-decoration:line-through underline;}}
.quick-table th:nth-child(1),.quick-table td:nth-child(1){{width:12%;}}
.quick-table th:nth-child(2),.quick-table td:nth-child(2){{width:8%;}}
.quick-table th:nth-child(3),.quick-table td:nth-child(3){{width:5%;}}
.quick-table th:nth-child(4),.quick-table td:nth-child(4){{width:22%;}}
.quick-table th:nth-child(5),.quick-table td:nth-child(5){{width:6%;}}
.quick-table th:nth-child(6),.quick-table td:nth-child(6){{width:47%;}}
.en-name{{color:#999;font-size:0.85em;font-weight:normal;}}
.variant-detail{{margin-bottom:8px;border:1px solid #333;border-radius:8px;overflow:hidden;}}
.variant-detail summary{{padding:12px 16px;background:#16213e;cursor:pointer;font-size:1.05em;user-select:none;}}
.variant-detail summary:hover{{background:#1a3a5e;}}
.variant-detail[open] summary{{border-bottom:1px solid #444;}}
.grouper-section{{margin:10px 20px;}}
.grouper-section h4{{color:#4ecca3;margin:10px 0 5px;font-size:1em;}}
.step-detail{{margin:4px 0 4px 20px;border-left:2px solid #333;}}
.step-detail summary{{padding:6px 12px;background:rgba(255,255,255,0.02);cursor:pointer;font-size:0.95em;display:flex;align-items:center;gap:8px;}}
.step-detail summary:hover{{background:rgba(233,69,96,0.06);}}
.step-label{{flex:1;}}
.step-type{{color:#888;font-size:0.85em;margin-left:8px;}}
.preview{{color:#f0a500;font-size:0.85em;margin-left:10px;}}
/* 对比按钮 */
.cmp-btn{{padding:2px 6px;border:1px solid #555;background:#0f3460;color:#ccc;cursor:pointer;border-radius:3px;font-size:0.85em;flex-shrink:0;white-space:nowrap;}}
.cmp-btn:hover{{border-color:#e94560;color:#fff;}}
.cmp-btn.added{{background:#e94560;border-color:#e94560;color:#fff;}}
/* 对比面板 */
.compare-panel{{position:fixed;bottom:0;left:0;right:0;background:#16213e;border-top:3px solid #e94560;z-index:200;max-height:45vh;overflow-y:auto;display:none;}}
.compare-panel.show{{display:block;}}
.compare-header{{padding:10px 16px;background:#0f3460;display:flex;justify-content:space-between;align-items:center;position:sticky;top:0;z-index:5;}}
.compare-header h3{{color:#e94560;font-size:1em;}}
.compare-header button{{padding:4px 12px;border:1px solid #555;background:transparent;color:#ccc;cursor:pointer;border-radius:3px;}}
.compare-header button:hover{{border-color:#e94560;}}
.compare-body{{display:flex;gap:0;overflow-x:auto;}}
.compare-col{{flex:1;min-width:300px;padding:8px;}}
.compare-col h4{{color:#4ecca3;font-size:0.9em;margin-bottom:6px;text-align:center;}}
.compare-col table{{width:100%;font-size:0.8em;border-collapse:collapse;}}
.compare-col td{{padding:3px 6px;border-bottom:1px solid #222;}}
.compare-col .diff{{background:rgba(233,69,96,0.2);}}
.compare-col .pn{{color:#bbb;}}
.compare-col .pv{{color:#4ecca3;font-family:Consolas,monospace;}}
/* 参数设置弹窗 */
.settings-overlay{{position:fixed;inset:0;background:rgba(0,0,0,0.7);z-index:500;display:none;justify-content:center;align-items:center;}}
.settings-overlay.show{{display:flex;}}
.settings-modal{{background:#1a1a2e;border:1px solid #e94560;border-radius:12px;padding:20px;max-width:600px;width:90%;max-height:80vh;overflow-y:auto;}}
.settings-modal h3{{color:#e94560;margin-bottom:12px;}}
.settings-modal .presets{{display:flex;gap:6px;margin-bottom:12px;flex-wrap:wrap;}}
.settings-modal .presets button{{padding:4px 12px;border:1px solid #555;background:#0f3460;color:#ccc;cursor:pointer;border-radius:3px;font-size:0.85em;}}
.settings-modal .presets button:hover{{border-color:#e94560;}}
.settings-modal .presets button.active{{background:#e94560;border-color:#e94560;}}
.param-grid{{display:grid;grid-template-columns:repeat(3,1fr);gap:4px;margin-bottom:12px;}}
.param-cb{{font-size:0.8em;color:#bbb;padding:2px 4px;display:flex;align-items:center;gap:4px;cursor:pointer;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;}}
.param-cb:hover{{color:#fff;}}
.settings-modal .btn-row{{display:flex;gap:8px;justify-content:flex-end;}}
.settings-modal .btn-row button{{padding:6px 16px;border:1px solid #555;background:#0f3460;color:#ccc;cursor:pointer;border-radius:4px;}}
.settings-modal .btn-row button.primary{{background:#e94560;border-color:#e94560;color:#fff;}}
.settings-modal .btn-row button:hover{{opacity:0.85;}}
/* 精简模式下行隐藏 — 仅隐藏有 data-param 的 tr，不动 thead */
.is-streamlined tr[data-param]:not(.mute-row){{display:none;}}
.is-streamlined tr.param-visible{{display:table-row;}}
/* 原有样式 */
.prop-table{{width:100%;table-layout:fixed;border-collapse:collapse;font-size:0.85em;margin:8px 0 12px 0;}}
.prop-table thead{{position:sticky;top:0;z-index:5;}}
.prop-table th{{background:#0a0a1a;padding:6px 10px;text-align:left;font-size:0.85em;color:#888;position:sticky;top:0;}}
.prop-table th:nth-child(1){{width:44%;}}
.prop-table th:nth-child(2){{width:28%;}}
.prop-table th:nth-child(3){{width:28%;}}
.prop-table td{{padding:4px 10px;border-bottom:1px solid #222;}}
.prop-table tbody tr:nth-child(even) td{{background:rgba(255,255,255,0.015);}}
.prop-table tbody tr:hover td{{background:rgba(233,69,96,0.06);}}
.prop-name{{color:#bbb;overflow:hidden;text-overflow:ellipsis;white-space:nowrap;}}
.prop-name:hover{{overflow:visible;white-space:normal;background:#1a1a2e;position:relative;z-index:5;}}
.prop-val{{color:#4ecca3;font-family:"Consolas","Courier New",monospace;white-space:nowrap;width:28%;}}
.prop-cat{{color:#666;font-size:0.8em;text-align:right;white-space:nowrap;width:28%;}}
.prop-cat.mod{{color:#f0a500;}}
.prop-cat.con{{color:#536dfe;}}
.mute-row{{display:none;}}
.show-mute .mute-row{{display:table-row;}}
.highlight{{background:#e9456040 !important;}}
.no-results{{text-align:center;padding:40px;color:#e94560;font-size:1.2em;display:none;}}
.back-to-top{{position:fixed;bottom:30px;right:30px;width:44px;height:44px;border-radius:50%;background:#e94560;color:#fff;border:none;cursor:pointer;font-size:1.2em;display:none;z-index:999;box-shadow:0 2px 12px rgba(233,69,96,0.4);transition:all 0.3s;opacity:0.85;}}
.back-to-top:hover{{opacity:1;transform:scale(1.1);}}
.back-to-top.show{{display:block;}}
.footer{{text-align:center;padding:30px;color:#666;font-size:0.85em;}}
.cmp-indicator{{position:fixed;top:12px;right:80px;background:#e94560;color:#fff;padding:6px 12px;border-radius:20px;font-size:0.85em;z-index:999;display:none;cursor:pointer;}}
.cmp-indicator.show{{display:block;}}
@media (max-width:768px){{
.tabs{{gap:2px;}}
.tab-btn{{padding:8px 10px;font-size:0.8em;}}
.content{{padding:10px;}}
.quick-table{{font-size:0.7em;}}
.search-bar input{{width:200px;}}
.search-bar{{flex-wrap:wrap;}}
.action-bar{{gap:4px;}}
.action-bar button{{padding:4px 8px;font-size:0.75em;}}
.back-to-top{{bottom:20px;right:20px;width:36px;height:36px;}}
.param-grid{{grid-template-columns:repeat(2,1fr);}}
.compare-col{{min-width:200px;}}
}}</style>
</head>
<body>

<div class="header">
<h1>🏔️ PEAK 地图参数手册</h1>
<p>全关卡 · 全变体 · 全生成器 · 全参数 — 中英双语 · 可搜索 · 可折叠</p>
<div class="stats">
<span class="stat"><strong>{len(segs)}</strong> 关卡</span>
<span class="stat"><strong>{total_variants}</strong> 变体</span>
<span class="stat"><strong>{total_steps}</strong> 生成步骤</span>
<span class="stat">更新 <strong>{now}</strong></span>
</div>
</div>

<div class="search-bar">
<span class="search-input-group">
<input type="text" id="searchInput" placeholder="🔍  搜索中文或英文 — 物件名、参数名、变体名、数值…" oninput="doSearch()" onkeydown="if(event.key==='Escape')clearSearch()">
<button class="search-clear" onclick="clearSearch()" title="清除搜索 (Esc)">✕</button>
</span>
<span class="search-info" id="searchInfo"></span>
</div>

<div class="tabs" id="tabBar">{''.join(tab_buttons)}</div>
<div class="no-results" id="noResults">😕 当前关卡没有匹配结果，试试切换到其他关卡标签</div>
<div class="cmp-indicator" id="cmpIndicator" onclick="showComparePanel()" title="点击查看对比">⊕ 对比 (0/2)</div>

<div class="content" id="contentArea">{''.join(tab_contents)}</div>
<button class="back-to-top" id="backToTop" onclick="scrollToTop()" title="回到顶部">⬆</button>

<!-- 对比面板 -->
<div class="compare-panel" id="comparePanel">
<div class="compare-header">
<h3>📊 生成器参数对比</h3>
<div><button onclick="clearCompare()">清空对比</button> <button onclick="closeComparePanel()">✕ 关闭</button></div>
</div>
<div class="compare-body" id="compareBody"></div>
</div>

<!-- 参数设置弹窗 -->
<div class="settings-overlay" id="settingsOverlay">
<div class="settings-modal">
<h3>⚙ 自定义参数显示</h3>
<p style="color:#888;font-size:0.85em;margin-bottom:10px;">勾选要显示的参数，未勾选的将被隐藏（精简模式）</p>
<div class="presets">
<button onclick="settingsPreset('core')" class="active">核心参数</button>
<button onclick="settingsPreset('all')">全部显示</button>
<button onclick="settingsPreset('none')">全部隐藏</button>
</div>
<div class="param-grid" id="paramGrid">
{param_checkboxes}
</div>
<div class="btn-row">
<button onclick="closeSettings()">取消</button>
<button class="primary" onclick="applySettings()">应用</button>
</div>
</div>
</div>

<div class="footer">
<p>数据来源：PEAK 1.62.a DreamyAscent Snapshot V2（27份诊断样本） | 翻译参考：TerrainCustomiserCN DisplayNameTranslator</p>
</div>

<script>
let currentTab=null,lastQuery='',currentSegId=null;
let compareList=[];
let paramSettings={{}};

function showTab(id){{
document.querySelectorAll('.tab-content').forEach(t=>t.classList.remove('active'));
document.querySelectorAll('.tab-btn').forEach(b=>b.classList.remove('active'));
const c=document.getElementById('tab-'+id);
if(c){{c.classList.add('active');currentTab=id;currentSegId=id;}}
document.querySelectorAll('.tab-btn').forEach(b=>{{if(b.textContent.includes(id.replace(/_/g,' '))||b.textContent.includes(id.replace(/_/g,'')))b.classList.add('active');}});
if(lastQuery)doSearch();
updateBackToTop();
applyCurrentSettings();
}}

function clearSearch(){{
document.getElementById('searchInput').value='';lastQuery='';
clearHighlights();
document.getElementById('searchInfo').textContent='';
document.getElementById('noResults').style.display='none';
}}

function clearHighlights(){{
document.querySelectorAll('.highlight').forEach(el=>el.classList.remove('highlight'));
}}

function toggleAll(segId,expand){{
const tc=document.getElementById('tab-'+segId);
if(!tc)return;
tc.querySelectorAll('.variant-detail,.step-detail').forEach(d=>{{d.open=expand;}});
}}

function toggleMuteRows(segId){{
const tc=document.getElementById('tab-'+segId);
if(!tc)return;
const btn=document.getElementById('btnMute_'+segId);
const isShowing=tc.classList.contains('show-mute');
if(isShowing){{tc.classList.remove('show-mute');if(btn){{btn.classList.remove('mute-visible');btn.innerHTML='🔇 显示静音参数';}}}}
else{{tc.classList.add('show-mute');if(btn){{btn.classList.add('mute-visible');btn.innerHTML='🔊 隐藏静音参数';}}}}
}}

function doSearch(){{
const q=document.getElementById('searchInput').value.trim();
lastQuery=q;clearHighlights();
document.getElementById('noResults').style.display='none';
document.getElementById('searchInfo').textContent='';
if(!q)return;
const ac=document.querySelector('.tab-content.active');
if(!ac)return;
const lq=q.toLowerCase();let fc=0;
ac.querySelectorAll('.variant-detail,.step-detail').forEach(d=>{{if(!d.open)d.open=true;}});
const fm={{el:null,top:Infinity}};
ac.querySelectorAll('tr,h4,summary,.seg-desc').forEach(row=>{{
if(row.textContent.toLowerCase().includes(lq)){{row.classList.add('highlight');fc++;
const r=row.getBoundingClientRect();if(r.top<fm.top&&r.top>0){{fm.el=row;fm.top=r.top;}}}}
}});
ac.querySelectorAll('.prop-val,.prop-name,.vn,.special').forEach(el=>{{
if(el.textContent.toLowerCase().includes(lq)){{el.classList.add('highlight');
const row=el.closest('tr');if(row&&!row.classList.contains('highlight')){{row.classList.add('highlight');fc++;}}}}
}});
if(fc===0){{document.getElementById('noResults').style.display='block';document.getElementById('searchInfo').textContent='0 个结果';}}
else{{document.getElementById('searchInfo').textContent='找到 '+fc+' 个结果';if(fm.el)fm.el.scrollIntoView({{behavior:'smooth',block:'center'}});}}
}}

function scrollToTop(){{window.scrollTo({{top:0,behavior:'smooth'}});}}
function updateBackToTop(){{
const btn=document.getElementById('backToTop');
if(!btn)return;
btn.classList.toggle('show',window.scrollY>400);
}}
window.addEventListener('scroll',updateBackToTop,{{passive:true}});

// ====== 对比功能 ======
function addToCompare(btn){{
const stepDetail=btn.closest('.step-detail');
if(!stepDetail)return;
const table=stepDetail.querySelector('.prop-table');
if(!table)return;
const summary=stepDetail.querySelector('summary');
const label=summary.querySelector('.step-label')?.textContent.trim()||'未知';
const variant=stepDetail.closest('.variant-detail')?.querySelector('summary strong')?.textContent.trim()||'';
const idx=compareList.findIndex(c=>c.btn===btn);
if(idx>=0){{compareList.splice(idx,1);btn.classList.remove('added');}}
else{{
if(compareList.length>=2){{const old=compareList.shift();if(old.btn)old.btn.classList.remove('added');}}
compareList.push({{btn,label,variant,tableHtml:table.outerHTML}});
btn.classList.add('added');
}}
updateCmpIndicator();
if(compareList.length===2)renderCompare();
}}

function updateCmpIndicator(){{
const ind=document.getElementById('cmpIndicator');
const n=compareList.length;
ind.textContent='⊕ 对比 ('+n+'/2)';
ind.classList.toggle('show',n>0);
}}

function showComparePanel(){{
document.getElementById('comparePanel').classList.add('show');
const btt=document.getElementById('backToTop');
if(btt)btt.style.bottom='50vh';
if(compareList.length>=1)renderCompare();
}}

function closeComparePanel(){{
document.getElementById('comparePanel').classList.remove('show');
const btt=document.getElementById('backToTop');
if(btt)btt.style.bottom='';
}}

function renderCompare(){{
const body=document.getElementById('compareBody');
if(compareList.length===0){{body.innerHTML='<div style="color:#888;text-align:center;padding:20px;">点击步骤旁的 ⊕ 按钮添加对比项</div>';return;}}
if(compareList.length===1){{
body.innerHTML='<div class="compare-col"><h4>'+escapeHtml(compareList[0].label)+' <small style="color:#888">('+escapeHtml(compareList[0].variant)+')</small></h4>'+compareList[0].tableHtml+'</div><div class="compare-col" style="color:#888;text-align:center;padding:40px;">再选一个步骤即可对比</div>';
return;
}}
const [a,b]=compareList;
const rowsA=parseRows(a.tableHtml);
const rowsB=parseRows(b.tableHtml);
const allKeys=new Set([...Object.keys(rowsA),...Object.keys(rowsB)]);
let html='';
html+='<div class="compare-col"><h4>'+escapeHtml(a.label)+' <small style="color:#888">('+escapeHtml(a.variant)+')</small></h4><table>';
for(const k of allKeys){{
const va=rowsA[k],vb=rowsB[k];
const diff=va!==undefined&&vb!==undefined&&va!==vb;
html+='<tr'+(diff?' class="diff"':'')+'><td class="pn">'+escapeHtml(k)+'</td><td class="pv">'+(va!==undefined?va:'—')+'</td></tr>';
}}
html+='</table></div>';
html+='<div class="compare-col"><h4>'+escapeHtml(b.label)+' <small style="color:#888">('+escapeHtml(b.variant)+')</small></h4><table>';
for(const k of allKeys){{
const va=rowsA[k],vb=rowsB[k];
const diff=va!==undefined&&vb!==undefined&&va!==vb;
html+='<tr'+(diff?' class="diff"':'')+'><td class="pn">'+escapeHtml(k)+'</td><td class="pv">'+(vb!==undefined?vb:'—')+'</td></tr>';
}}
html+='</table></div>';
body.innerHTML=html;
document.getElementById('comparePanel').classList.add('show');
}}

function parseRows(tableHtml){{
const div=document.createElement('div');div.innerHTML=tableHtml;
const rows={{}};
div.querySelectorAll('tr').forEach(tr=>{{
const nameEl=tr.querySelector('.prop-name');
const valEl=tr.querySelector('.prop-val');
if(nameEl&&valEl)rows[nameEl.textContent.trim()]=valEl.textContent.trim();
}});
return rows;
}}

function clearCompare(){{
compareList.forEach(c=>{{if(c.btn)c.btn.classList.remove('added');}});
compareList=[];
updateCmpIndicator();
closeComparePanel();
}}

// ====== 参数设置功能 ======
function openSettings(segId){{
currentSegId=segId;
document.getElementById('settingsOverlay').classList.add('show');
const cbs=document.querySelectorAll('#paramGrid input[type=checkbox]');
const hasSettings=Object.keys(paramSettings).length>0;
if(!hasSettings){{
const core=['area','nrOfSpawns','chanceToUseSpawner','layerType'];
cbs.forEach(cb=>{{cb.checked=core.includes(cb.value);}});
}}else{{
cbs.forEach(cb=>{{if(paramSettings.hasOwnProperty(cb.value))cb.checked=paramSettings[cb.value];}});
}}
document.querySelectorAll('.presets button').forEach(b=>b.classList.remove('active'));
const arr=Array.from(cbs);
if(arr.every(cb=>cb.checked))document.querySelector('.presets button:nth-child(2)').classList.add('active');
else if(arr.every(cb=>!cb.checked))document.querySelector('.presets button:nth-child(3)').classList.add('active');
else document.querySelector('.presets button:nth-child(1)').classList.add('active');
}}

function closeSettings(){{document.getElementById('settingsOverlay').classList.remove('show');}}

function settingsPreset(type){{
const cbs=document.querySelectorAll('#paramGrid input[type=checkbox]');
document.querySelectorAll('.presets button').forEach(b=>b.classList.remove('active'));
if(type==='core'){{
const core=['area','nrOfSpawns','chanceToUseSpawner','layerType'];
cbs.forEach(cb=>{{cb.checked=core.includes(cb.value);}});
document.querySelector('.presets button:nth-child(1)').classList.add('active');
}}else if(type==='all'){{
cbs.forEach(cb=>{{cb.checked=true;}});
document.querySelector('.presets button:nth-child(2)').classList.add('active');
}}else{{
cbs.forEach(cb=>{{cb.checked=false;}});
document.querySelector('.presets button:nth-child(3)').classList.add('active');
}}
}}

function applySettings(){{
const cbs=document.querySelectorAll('#paramGrid input[type=checkbox]');
paramSettings={{}};
cbs.forEach(cb=>{{paramSettings[cb.value]=cb.checked;}});
applyCurrentSettings();
closeSettings();
}}

function applyCurrentSettings(){{
if(!currentSegId)return;
const tc=document.getElementById('tab-'+currentSegId);
if(!tc)return;
const keys=Object.keys(paramSettings);
if(keys.length===0){{tc.classList.remove('is-streamlined');return;}}
const allOn=keys.every(k=>paramSettings[k]);
if(allOn){{tc.classList.remove('is-streamlined');return;}}
tc.classList.add('is-streamlined');
tc.querySelectorAll('tr[data-param]').forEach(tr=>{{
const pk=tr.getAttribute('data-param');
if(paramSettings[pk])tr.classList.add('param-visible');
else tr.classList.remove('param-visible');
}});
}}

function escapeHtml(s){{return String(s).replace(/&/g,'&amp;').replace(/</g,'&lt;').replace(/>/g,'&gt;').replace(/"/g,'&quot;');}}

function jumpStep(e, anchor){{
e.preventDefault();
const el=document.getElementById(anchor);
if(!el)return;
let p=el.parentElement;
while(p){{
if(p.tagName==='DETAILS')p.open=true;
p=p.parentElement;
}}
el.scrollIntoView({{behavior:'smooth',block:'center'}});
el.style.background='rgba(233,69,96,0.12)';
setTimeout(()=>el.style.background='',1800);
window.location.hash=anchor;
}}

document.addEventListener('DOMContentLoaded',()=>{{
const fb=document.querySelector('.tab-btn');
if(fb)fb.click();
updateBackToTop();
}});

document.addEventListener('keydown',(e)=>{{
if((e.ctrlKey||e.metaKey)&&e.key==='f'){{e.preventDefault();document.getElementById('searchInput').focus();document.getElementById('searchInput').select();}}
if(e.key==='Escape'){{closeSettings();closeComparePanel();}}
}});
</script>

</body>
</html>"""

    html_path=OUT_DIR/"PEAK地图参数手册.html"
    with open(str(html_path),"w",encoding="utf-8") as f:
        f.write(full)
    size_kb=html_path.stat().st_size/1024
    print(f"  ✓ HTML: {html_path.name} ({size_kb:.0f} KB)")


def gen_excel(data):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("  ⚠ openpyxl 未安装,跳过 Excel")
        return

    wb=Workbook();wb.remove(wb.active)
    hf=Font(name="Microsoft YaHei",bold=True,color="FFFFFF",size=11)
    hfl=PatternFill(start_color="0F3460",end_color="0F3460",fill_type="solid")
    nf=Font(name="Microsoft YaHei",size=10)
    tb=Border(left=Side(style="thin",color="444444"),right=Side(style="thin",color="444444"),top=Side(style="thin",color="444444"),bottom=Side(style="thin",color="444444"))

    segs=data["segments"]

    ws=wb.create_sheet("总览")
    for ci,h in enumerate(["关卡序号","中文名","英文名","变体数","分组器数","步骤数","简述"],1):
        c=ws.cell(row=1,column=ci,value=h);c.font=hf;c.fill=hfl;c.border=tb
    row=2
    for sn in SEGMENT_ORDER:
        if sn not in segs: continue
        seg=segs[sn];info=SEGMENT_INFO.get(sn,{})
        vn=len(seg["variants"]);gn=sum(len(v["groupers"]) for v in seg["variants"].values());sn_count=sum(sum(len(g["steps"]) for g in v["groupers"]) for v in seg["variants"].values())
        for ci,v in enumerate([info.get("order",""),seg["nameCN"],sn,vn,gn,sn_count,info.get("desc","")],1):
            c=ws.cell(row=row,column=ci,value=v);c.font=nf;c.border=tb
        row+=1
    ws.column_dimensions["A"].width=10;ws.column_dimensions["B"].width=16;ws.column_dimensions["C"].width=18
    ws.column_dimensions["D"].width=8;ws.column_dimensions["E"].width=10;ws.column_dimensions["F"].width=8;ws.column_dimensions["G"].width=65
    ws.freeze_panes="A2"

    for sn in SEGMENT_ORDER:
        if sn not in segs: continue
        seg=segs[sn];scn=seg["nameCN"][:31]
        ws=wb.create_sheet(scn)
        hds=["变体 (Variant)","分组器 (Grouper)","步骤·中文 (Step CN)","步骤·英文 (Step EN)","类型 (Type)","参数·英文 (Param EN)","参数值 (Value)","来源 (Source)","来源名 (Source Name)"]
        for ci,h in enumerate(hds,1):
            c=ws.cell(row=1,column=ci,value=h);c.font=hf;c.fill=hfl;c.border=tb
        row=2
        for vn,v in seg["variants"].items():
            vnf=fix_variant_name(vn);vc=v.get("nameCN",vnf)
            if vc=="None" or vc is None: vc=vnf
            vl=f"{vc} ({vnf})" if vc!=vnf else vc
            for g in v["groupers"]:
                gl=f"{g['nameCN']} ({g['name']})" if g['nameCN']!=g['name'] else g['name']
                for st in g["steps"]:
                    base=[vl,gl,st["nameCN"],st["name"],f"{st.get('typeCN',st['type'])} ({st['type']})"]
                    for pk,pv in st.get("properties",{}).items():
                        for ci,v in enumerate(base+[pk,format_val_short(pv),"属性",""],1):
                            c=ws.cell(row=row,column=ci,value=v);c.font=nf;c.border=tb
                        row+=1
                    for mn,mps in st.get("modifiers",{}).items():
                        for pk,pv in mps.items():
                            for ci,v in enumerate(base+[pk,format_val_short(pv),"修改器",mn],1):
                                c=ws.cell(row=row,column=ci,value=v);c.font=nf;c.border=tb
                            row+=1
                    for cn,cps in st.get("constraints",{}).items():
                        for pk,pv in cps.items():
                            for ci,v in enumerate(base+[pk,format_val_short(pv),"条件",cn],1):
                                c=ws.cell(row=row,column=ci,value=v);c.font=nf;c.border=tb
                            row+=1
        for ci,w in enumerate([22,22,22,20,20,18,14,12,22],1):
            ws.column_dimensions[get_column_letter(ci)].width=w
        ws.freeze_panes="A2";ws.auto_filter.ref=f"A1:{get_column_letter(len(hds))}{row-1}"

    xp=OUT_DIR/"PEAK地图参数速查.xlsx";wb.save(str(xp))
    print(f"  ✓ Excel: {xp.name} ({xp.stat().st_size/1024:.0f} KB)")


def main():
    print("\n"+"="*60)
    print("PEAK 地图参数文档生成 (v3: +对比 +精简)")
    print("="*60)
    data=load_data()
    gen_html(data)
    gen_excel(data)
    print("\n完成! 输出:",str(OUT_DIR))

if __name__=="__main__":
    main()
